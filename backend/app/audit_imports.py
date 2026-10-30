"""Parsing and deterministic matching for physical-audit workbook evidence."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from numbers import Number
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Asset, AuditMatchReason, AuditMatchStrategy
from .system_imports import (
    ImportValidationError,
    _json_value,
    _text,
    _validate_xlsx_archive,
    normalize_asset_code,
    original_asset_code,
)

AUDIT_HEADERS = (
    "CANTIDAD",
    "IDENTIFICACION",
    "MAQUINA/EQUIPO",
    "MARCA",
    "MODELO ",
    "FUNCIONAMIENTO",
    "ESTADO EXTERNO",
    "OBSERVACIONES",
    "EXISTENCIA EN BRIDGE",
    "CC FISICO",
    "CC  BG",
)

BG_AUDIT_HEADERS = {
    "Activos en sistema BG Pag. 1": (
        "RUBRO", "CATEGORIA", "PRODUCTO", "CODIGO DE INDENTIFICACION", "IDENTIFICACION", "NRO. DEPOSITO", "DEPOSITO", "ESTADO",
    ),
    "Activos en sistema BG Pag. 2": (
        "RUBRO", "CATEGORIA", "PRODUCTO", "CODIGO DE IDENTIFICACION", "IDENTIFICACION", "NRO. DEPOSITO", "DEPOSITO", "ESTADO",
    ),
}
BG_AUDIT_SHEETS = tuple(BG_AUDIT_HEADERS)
BG_IDENTIFIER_COLUMN = 4
BG_DIRECT_FILL_SIGNATURE_KEY = "bg_direct_fill_signature"
MANUAL_DIRECT_FILL_SIGNATURE_KEY = "manual_direct_fill_signature"
BG_RETURN_FILL_SIGNATURE = {
    "fill_type": "solid",
    "foreground": {"type": "theme", "theme": 7, "tint": 0.5999938962981048},
}


@dataclass(frozen=True)
class ParsedAuditRow:
    """One physical-audit row and the complete positional source evidence."""

    named_values: dict[str, Any]
    original_data: dict[str, Any]
    quantity: Decimal
    identifier: str
    source_kind: str


@dataclass(frozen=True)
class ParsedAuditReport:
    rows: tuple[ParsedAuditRow, ...]
    sheet_names: tuple[str, ...]
    bg_row_count: int = 0
    manual_row_count: int = 0
    is_mixed_audit: bool = False


@dataclass(frozen=True)
class AuditMatch:
    candidate_asset: Asset | None
    strategy: AuditMatchStrategy | None
    reason: AuditMatchReason | None


def _is_blank_row(values: tuple[Any, ...]) -> bool:
    return not any(_text(value) for value in values)


def _quantity(value: Any) -> Decimal:
    """Accept only source numeric, positive whole-unit audit quantities."""
    if isinstance(value, bool) or not isinstance(value, Number):
        raise ImportValidationError("invalid_quantity", "every CANTIDAD value must be a positive integer")
    quantity = Decimal(str(value))
    if not quantity.is_finite() or quantity <= 0 or quantity != quantity.to_integral_value():
        raise ImportValidationError("invalid_quantity", "every CANTIDAD value must be a positive integer")
    return quantity


def _color_signature(color: Any) -> dict[str, Any]:
    """Render the stored color fields without resolving theme/table/CF styles."""
    signature: dict[str, Any] = {"type": color.type, "tint": color.tint}
    if color.type == "theme":
        signature["theme"] = color.theme
    elif color.type == "rgb":
        signature["rgb"] = color.rgb
    elif color.type == "indexed":
        signature["indexed"] = color.indexed
    elif color.type == "auto":
        signature["auto"] = color.auto
    return signature


def _direct_fill_signature(cell: Any) -> dict[str, Any]:
    """Capture only the direct fill stored on a cell, without style inference."""
    return {
        "fill_type": cell.fill.fill_type,
        "foreground": _color_signature(cell.fill.fgColor),
        "background": _color_signature(cell.fill.bgColor),
    }


def _is_bg_return_fill_signature(signature: list[dict[str, Any]]) -> bool:
    """Classify only the explicitly confirmed direct-fill attributes."""
    return len(signature) == 8 and all(
        cell.get("fill_type") == BG_RETURN_FILL_SIGNATURE["fill_type"]
        and cell.get("foreground") == BG_RETURN_FILL_SIGNATURE["foreground"]
        for cell in signature
    )


def _parse_bg_audit_report(workbook: Any) -> ParsedAuditReport:
    """Parse the complete, exact BG page pair and preserve row-fill evidence."""
    present_sheets = [name for name in BG_AUDIT_SHEETS if name in workbook.sheetnames]
    if present_sheets and len(present_sheets) != len(BG_AUDIT_SHEETS):
        raise ImportValidationError(
            "incomplete_bg_audit_profile",
            "BG audit imports require both Activos en sistema BG page sheets",
        )
    if not present_sheets:
        raise ImportValidationError("no_bg_audit_sheets", "workbook contains no BG audit sheets")

    rows: list[ParsedAuditRow] = []
    for sheet_name in BG_AUDIT_SHEETS:
        worksheet = workbook[sheet_name]
        header_cells = tuple(next(worksheet.iter_rows(min_row=1, max_row=1), ()) or ())
        header_values = tuple(cell.value for cell in header_cells)
        positioned = tuple("" if index >= len(header_values) or header_values[index] is None else str(header_values[index]) for index in range(8))
        if positioned != BG_AUDIT_HEADERS[sheet_name]:
            raise ImportValidationError(
                "invalid_bg_audit_layout",
                "each BG audit page must contain its exact A:H headers on row 1",
            )
        for source_row, source_cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
            values = tuple(cell.value for cell in source_cells)
            if _is_blank_row(values):
                continue
            cells = source_cells[:8]
            original_data = {
                "source_sheet": worksheet.title,
                "source_row": source_row,
                "cells": [
                    {
                        "column": index + 1,
                        "header": _json_value(header_values[index]) if index < len(header_values) else None,
                        "value": _json_value(cell.value),
                    }
                    for index, cell in enumerate(source_cells)
                ],
                BG_DIRECT_FILL_SIGNATURE_KEY: [_direct_fill_signature(cell) for cell in cells],
            }
            named_values = {
                str(header_values[index]): _json_value(cell.value)
                for index, cell in enumerate(cells)
            }
            named_values["FUNCIONAMIENTO"] = _json_value(cells[7].value)
            rows.append(
                ParsedAuditRow(
                    named_values=named_values,
                    original_data=original_data,
                    quantity=Decimal("1"),
                    identifier=original_asset_code(cells[BG_IDENTIFIER_COLUMN - 1].value),
                    source_kind="bg",
                )
            )
    if not rows:
        raise ImportValidationError("no_usable_rows", "workbook contains no usable BG audit rows")
    return ParsedAuditReport(rows=tuple(rows), sheet_names=BG_AUDIT_SHEETS, bg_row_count=len(rows))


def _parse_manual_audit_sheet(worksheet: Any) -> tuple[ParsedAuditRow, ...] | None:
    """Parse one row-three A:K physical sheet without using its display name."""
    header_cells = tuple(next(worksheet.iter_rows(min_row=3, max_row=3), ()) or ())
    header_values = tuple(cell.value for cell in header_cells)
    headers = tuple("" if value is None else str(value) for value in header_values)
    present = [header for header in AUDIT_HEADERS if header in headers]
    if not present:
        return None
    missing = [header for header in AUDIT_HEADERS if header not in headers]
    duplicates = [header for header in AUDIT_HEADERS if headers.count(header) != 1]
    positioned = tuple(headers[index] if index < len(headers) else "" for index in range(len(AUDIT_HEADERS)))
    if missing or duplicates or positioned != AUDIT_HEADERS:
        raise ImportValidationError(
            "invalid_audit_layout",
            "each audit sheet must contain the required physical-audit headers exactly once in A:K on row 3",
        )
    indexes = {header: headers.index(header) for header in AUDIT_HEADERS}
    rows: list[ParsedAuditRow] = []
    for source_row, source_cells in enumerate(worksheet.iter_rows(min_row=4), start=4):
        values = tuple(cell.value for cell in source_cells)
        # Formatting alone is not source evidence and must never create a row.
        if _is_blank_row(values):
            continue
        cells = [
            {
                "column": index + 1,
                "header": _json_value(header_values[index]) if index < len(header_values) else None,
                "value": _json_value(cell.value),
            }
            for index, cell in enumerate(source_cells)
        ]
        named_values = {
            header: _json_value(values[index] if index < len(values) else None)
            for header, index in indexes.items()
        }
        rows.append(
            ParsedAuditRow(
                named_values=named_values,
                original_data={
                    "source_sheet": worksheet.title,
                    "source_row": source_row,
                    "cells": cells,
                    MANUAL_DIRECT_FILL_SIGNATURE_KEY: [_direct_fill_signature(cell) for cell in source_cells],
                },
                quantity=_quantity(named_values["CANTIDAD"]),
                identifier=original_asset_code(named_values["IDENTIFICACION"]),
                source_kind="manual",
            )
        )
    return tuple(rows)


def parse_audit_report(content: bytes) -> ParsedAuditReport:
    """Read compatible physical sheets, including the strict mixed BG/manual profile."""
    _validate_xlsx_archive(content)
    workbook = None
    try:
        # Direct fill metadata is source evidence for both profile types.
        workbook = load_workbook(filename=BytesIO(content), read_only=False, data_only=False)
        has_bg_profile = any(name in workbook.sheetnames for name in BG_AUDIT_SHEETS)
        if has_bg_profile:
            bg_report = _parse_bg_audit_report(workbook)
            manual_sheets: list[str] = []
            manual_rows: list[ParsedAuditRow] = []
            for worksheet in workbook.worksheets:
                if worksheet.title in BG_AUDIT_SHEETS:
                    continue
                parsed_rows = _parse_manual_audit_sheet(worksheet)
                if parsed_rows is None:
                    continue
                manual_sheets.append(worksheet.title)
                manual_rows.extend(parsed_rows)
            if not manual_sheets:
                return bg_report
            if len(manual_sheets) != 5:
                raise ImportValidationError(
                    "invalid_mixed_audit_profile",
                    "mixed BG audit imports require exactly five physical-audit sheets using the A:K row-three layout",
                )
            return ParsedAuditReport(
                rows=(*bg_report.rows, *manual_rows),
                sheet_names=(*BG_AUDIT_SHEETS, *manual_sheets),
                bg_row_count=bg_report.bg_row_count,
                manual_row_count=len(manual_rows),
                is_mixed_audit=True,
            )

        rows: list[ParsedAuditRow] = []
        audit_sheets: list[str] = []
        for worksheet in workbook.worksheets:
            parsed_rows = _parse_manual_audit_sheet(worksheet)
            if parsed_rows is None:
                continue
            audit_sheets.append(worksheet.title)
            rows.extend(parsed_rows)
        if not audit_sheets:
            raise ImportValidationError("no_audit_sheets", "workbook contains no physical-audit sheets")
        if not rows:
            raise ImportValidationError("no_usable_rows", "workbook contains no usable audit rows")
        return ParsedAuditReport(rows=tuple(rows), sheet_names=tuple(audit_sheets), manual_row_count=len(rows))
    except ImportValidationError:
        raise
    except Exception as error:
        raise ImportValidationError("unreadable_workbook", "file is not a readable .xlsx workbook") from error
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def match_audit_identifier(db: Session, identifier: str) -> AuditMatch:
    """Match only exact original code, then one unique normalized code."""
    if not identifier.strip():
        return AuditMatch(None, None, AuditMatchReason.MISSING_IDENTIFIER)
    exact = db.scalar(select(Asset).where(Asset.original_code == identifier))
    if exact is not None:
        return AuditMatch(exact, AuditMatchStrategy.EXACT_ORIGINAL_CODE, None)
    candidates = list(db.scalars(select(Asset).where(Asset.normalized_code == normalize_asset_code(identifier))))
    if not candidates:
        return AuditMatch(None, None, AuditMatchReason.NO_NORMALIZED_CANDIDATE)
    if len(candidates) != 1:
        return AuditMatch(None, None, AuditMatchReason.AMBIGUOUS_NORMALIZED_CANDIDATE)
    return AuditMatch(candidates[0], AuditMatchStrategy.NORMALIZED_CODE, None)
