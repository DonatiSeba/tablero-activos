"""Parsing and immutable local storage for CC190-style system snapshots."""

from __future__ import annotations

import hashlib
import os
from io import BytesIO
import re
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

REQUIRED_COLUMNS = (
    "Rubro",
    "Categoría",
    "Producto",
    "Cód. Ident.",
    "Identificación",
    "Nro. CC",
    "Centro de Costo",
    "Estado",
)
DEFAULT_MAX_UPLOAD_BYTES = 25 * 1024 * 1024
DEFAULT_MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
UPLOAD_READ_CHUNK_SIZE = 1024 * 1024


class ImportValidationError(ValueError):
    """A client-safe reason why a system report cannot become evidence."""

    def __init__(self, code: str, detail: str) -> None:
        super().__init__(detail)
        self.code = code
        self.detail = detail


class StorageIntegrityError(OSError):
    """A digest-addressed evidence path exists but does not contain its digest."""


@dataclass(frozen=True)
class ParsedSystemRow:
    """One source row with named parsing fields and collision-free raw evidence."""

    named_values: dict[str, Any]
    original_data: dict[str, list[dict[str, Any]]]


@dataclass(frozen=True)
class ParsedSystemReport:
    """Validated source rows and the one cost center asserted by a report."""

    rows: tuple[ParsedSystemRow, ...]
    cost_center_code: str
    observed_cost_center_names: tuple[str, ...]


def _configured_positive_limit(name: str, default: int) -> int:
    try:
        value = int(os.environ.get(name, str(default)))
    except ValueError as error:
        raise RuntimeError(f"{name} must be a positive integer") from error
    if value <= 0:
        raise RuntimeError(f"{name} must be a positive integer")
    return value


def max_upload_bytes() -> int:
    """Return the maximum accepted upload size before workbook processing."""
    return _configured_positive_limit("IMPORT_MAX_UPLOAD_BYTES", DEFAULT_MAX_UPLOAD_BYTES)


def max_xlsx_uncompressed_bytes() -> int:
    """Return the maximum declared XLSX archive payload size."""
    return _configured_positive_limit(
        "IMPORT_MAX_XLSX_UNCOMPRESSED_BYTES", DEFAULT_MAX_XLSX_UNCOMPRESSED_BYTES
    )


def read_upload_content(stream: BinaryIO) -> bytes:
    """Read an upload in bounded chunks without persisting unvalidated bytes."""
    limit = max_upload_bytes()
    content = bytearray()
    while True:
        chunk = stream.read(UPLOAD_READ_CHUNK_SIZE)
        if not chunk:
            break
        if len(content) + len(chunk) > limit:
            content.clear()
            raise ImportValidationError("upload_too_large", "uploaded file exceeds the configured size limit")
        content.extend(chunk)
    return bytes(content)


def parse_report_date(value: str | None) -> date:
    """Accept only the calendar-date ISO representation used by the API."""
    if value is None or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ImportValidationError("invalid_report_date", "report_date must be an ISO date (YYYY-MM-DD)")
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise ImportValidationError("invalid_report_date", "report_date must be an ISO date (YYYY-MM-DD)") from error


def _json_value(value: Any) -> Any:
    """Preserve worksheet cell values in JSON without applying domain normalization."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return str(value)


def _text(value: Any) -> str:
    """Render a report identifier deterministically while retaining the raw cell elsewhere."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def original_asset_code(value: Any) -> str:
    """Preserve a textual asset code exactly; only numeric Excel rendering is canonicalized."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize_asset_code(original_code: str) -> str:
    """Apply only deterministic presentation normalization, never identity matching."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", original_code).upper())


def _validate_xlsx_archive(content: bytes) -> None:
    """Reject malformed archives and archives exceeding the configured expanded size."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > max_xlsx_uncompressed_bytes():
                raise ImportValidationError(
                    "xlsx_uncompressed_size_exceeded",
                    "workbook exceeds the configured uncompressed size limit",
                )
    except ImportValidationError:
        raise
    except Exception as error:
        raise ImportValidationError("unreadable_workbook", "file is not a readable .xlsx workbook") from error


def parse_system_report(content: bytes) -> ParsedSystemReport:
    """Read the first worksheet and validate the report's named-column contract."""
    _validate_xlsx_archive(content)
    workbook = None
    try:
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=False)
        worksheet = workbook.active
        header_values = tuple(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()) or ())
        headers = tuple("" if value is None else str(value) for value in header_values)
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        duplicates = [column for column in REQUIRED_COLUMNS if headers.count(column) != 1]
        if missing or duplicates:
            raise ImportValidationError(
                "missing_required_columns",
                "workbook must contain each required column exactly once: " + ", ".join(REQUIRED_COLUMNS),
            )
        indexes = {column: headers.index(column) for column in REQUIRED_COLUMNS}
        rows: list[ParsedSystemRow] = []
        cost_center_codes: set[str] = set()
        cost_center_names: set[str] = set()
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            # Preserve cells by position rather than source header as a JSON key:
            # blank, duplicate, and generated-looking headers cannot collide.
            column_count = max(len(header_values), len(values))
            cells = [
                {
                    "column": index + 1,
                    "header": _json_value(header_values[index]) if index < len(header_values) else None,
                    "value": _json_value(values[index]) if index < len(values) else None,
                }
                for index in range(column_count)
            ]
            if not any(_text(cell["value"]) for cell in cells):
                continue
            named_values = {
                column: _json_value(values[index] if index < len(values) else None)
                for column, index in indexes.items()
            }
            code = _text(named_values["Nro. CC"])
            if not code:
                raise ImportValidationError("blank_cost_center_code", "every non-blank row must contain Nro. CC")
            cost_center_codes.add(code)
            name = _text(named_values["Centro de Costo"])
            if name:
                cost_center_names.add(name)
            rows.append(ParsedSystemRow(named_values=named_values, original_data={"cells": cells}))
        if not rows:
            raise ImportValidationError("no_usable_rows", "workbook contains no usable data rows")
        if len(cost_center_codes) != 1:
            raise ImportValidationError("nonuniform_cost_center_code", "workbook must contain exactly one Nro. CC value")
        workbook.close()
        workbook = None
        return ParsedSystemReport(
            rows=tuple(rows),
            cost_center_code=next(iter(cost_center_codes)),
            observed_cost_center_names=tuple(sorted(cost_center_names)),
        )
    except ImportValidationError:
        raise
    except Exception as error:
        raise ImportValidationError("unreadable_workbook", "file is not a readable .xlsx workbook") from error
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                # The original malformed-workbook failure is the client-safe rejection.
                pass


def sha256_hex(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def import_storage_path() -> Path:
    """Return the configured local evidence root without deriving paths from input names."""
    return Path(os.environ.get("IMPORT_STORAGE_PATH", "/data/imports")).expanduser()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stored:
        for chunk in iter(lambda: stored.read(UPLOAD_READ_CHUNK_SIZE), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _verify_existing_file(target: Path, sha256: str) -> None:
    if _sha256_file(target) != sha256:
        raise StorageIntegrityError("existing digest-addressed evidence is corrupt")


def _fsync_directory(directory: Path) -> None:
    """Persist a directory entry where the host operating system supports it."""
    if os.name == "nt":
        # Windows does not permit opening a directory with os.open for fsync.
        return
    descriptor = os.open(directory, os.O_RDONLY)
    try:
        os.fsync(descriptor)
    finally:
        os.close(descriptor)


def store_original_file(content: bytes, sha256: str) -> tuple[str, bool]:
    """Atomically publish verified content without replacing an existing digest path."""
    root = import_storage_path()
    target = root / sha256[:2] / f"{sha256}.xlsx"
    target.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    relative_path = str(target.relative_to(root))
    if target.exists():
        _verify_existing_file(target, sha256)
        return relative_path, False

    temporary = target.parent / f".{sha256}.{uuid.uuid4().hex}.tmp"
    descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(descriptor, "wb") as output:
            output.write(content)
            output.flush()
            os.fsync(output.fileno())
        try:
            os.link(temporary, target)
        except FileExistsError:
            _verify_existing_file(target, sha256)
            return relative_path, False
        _fsync_directory(target.parent)
        return relative_path, True
    finally:
        temporary.unlink(missing_ok=True)
