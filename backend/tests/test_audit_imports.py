import asyncio
from collections.abc import Generator
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import imports as imports_api
from backend.app.audit_imports import AUDIT_HEADERS, match_audit_identifier, parse_audit_report
from backend.app.auth import password_hasher
from backend.app.db import Base, get_db
from backend.app.main import app
from backend.app.models import (
    Asset,
    AssetAlias,
    AssetObservation,
    AuditCurrentState,
    AuditLog,
    AuditCurrentStateProjection,
    AuditMatchReason,
    AuditMatchStrategy,
    CostCenter,
    CostCenterStatus,
    ImportBatch,
    ReconciliationCase,
    User,
    UserRole,
)
from backend.app.system_imports import ImportValidationError, normalize_asset_code


def workbook_bytes(
    rows_by_sheet: dict[str, list[list[object]]], headers: list[str] | None = None, reference_sheet: bool = False
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in rows_by_sheet.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Physical audit"])
        sheet.append([])
        sheet.append(headers or list(AUDIT_HEADERS))
        for row in rows:
            sheet.append(row)
    if reference_sheet:
        reference = workbook.create_sheet("Activos en sistema BG reference")
        reference.append(["Bridge reference only"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def database(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Generator[sessionmaker[Session], None, None]:
    monkeypatch.setenv("SESSION_SECRET", "test-session-secret-that-is-long-enough-for-signing")
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("SESSION_TTL_SECONDS", "3600")
    monkeypatch.setenv("IMPORT_STORAGE_PATH", str(tmp_path / "imports"))
    engine = create_engine("sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def override_get_db() -> Generator[Session, None, None]:
        db = factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    yield factory
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()


def add_user_and_center(factory: sessionmaker[Session], role: UserRole = UserRole.EDITOR) -> None:
    with factory() as db:
        db.add_all(
            [
                User(
                    username="audit-user", email="audit@example.test", display_name="Audit User",
                    password_hash=password_hasher.hash("correct horse battery staple"), role=role, is_active=True,
                ),
                CostCenter(code="190", name="Operations", status=CostCenterStatus.ACTIVE),
            ]
        )
        db.commit()


async def client() -> httpx.AsyncClient:
    result = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver")
    login = await result.post("/api/auth/login", json={"username": "audit-user", "password": "correct horse battery staple"})
    assert login.status_code == 200
    return result


def test_audit_parser_requires_row_three_a_to_k_and_preserves_l_and_source_position() -> None:
    content = workbook_bytes(
        {"Physical 1": [[2, " Asset  01 ", "machine", "brand", "model", "works", "good", "private note", "yes", "190", "190", ""]]},
        reference_sheet=True,
    )
    report = parse_audit_report(content)
    row = report.rows[0]
    assert report.sheet_names == ("Physical 1",)
    assert row.quantity == 2
    assert row.identifier == " Asset  01 "
    assert row.original_data["source_sheet"] == "Physical 1"
    assert row.original_data["source_row"] == 4
    assert row.original_data["cells"][11] == {"column": 12, "header": None, "value": None}

    malformed_headers = list(AUDIT_HEADERS)
    malformed_headers[4] = "MODELO"
    with pytest.raises(ImportValidationError) as malformed:
        parse_audit_report(workbook_bytes({"Attempted audit": [[1] * 11]}, malformed_headers))
    assert malformed.value.code == "invalid_audit_layout"
    with pytest.raises(ImportValidationError) as duplicate:
        parse_audit_report(workbook_bytes({"Attempted audit": [[1] * 12]}, list(AUDIT_HEADERS) + ["CANTIDAD"]))
    assert duplicate.value.code == "invalid_audit_layout"
    with pytest.raises(ImportValidationError) as absent:
        parse_audit_report(workbook_bytes({}, reference_sheet=True))
    assert absent.value.code == "no_audit_sheets"


@pytest.mark.parametrize("quantity", [0, -1, 1.5, "1", True, None])
def test_audit_parser_rejects_non_positive_or_non_numeric_quantities(quantity: object) -> None:
    with pytest.raises(ImportValidationError) as invalid:
        parse_audit_report(workbook_bytes({"Physical": [[quantity, None, "machine"] + [None] * 8]}))
    assert invalid.value.code == "invalid_quantity"


def test_audit_parser_skips_fully_blank_tail_rows() -> None:
    report = parse_audit_report(
        workbook_bytes({"Physical": [[1, "AUDIT-01", "machine"] + [None] * 8, [None] * 11]})
    )
    assert len(report.rows) == 1
    assert report.rows[0].original_data["source_row"] == 4


def test_matching_is_exact_then_unique_normalized_only(database: sessionmaker[Session]) -> None:
    with database() as db:
        exact = Asset(original_code="A B", normalized_code=normalize_asset_code("A B"), description="do not search description")
        normalized = Asset(original_code="C D", normalized_code=normalize_asset_code("C D"))
        duplicate_one = Asset(original_code="D E", normalized_code=normalize_asset_code("D E"))
        duplicate_two = Asset(original_code="DE", normalized_code=normalize_asset_code("D E"))
        db.add_all([exact, normalized, duplicate_one, duplicate_two])
        db.commit()
        assert match_audit_identifier(db, "A B").strategy is AuditMatchStrategy.EXACT_ORIGINAL_CODE
        unique = match_audit_identifier(db, " C  D ")
        assert unique.candidate_asset is normalized
        assert unique.strategy is AuditMatchStrategy.NORMALIZED_CODE
        assert match_audit_identifier(db, " D E ").reason is AuditMatchReason.AMBIGUOUS_NORMALIZED_CANDIDATE
        assert match_audit_identifier(db, "not present").reason is AuditMatchReason.NO_NORMALIZED_CANDIDATE
        assert match_audit_identifier(db, " ").reason is AuditMatchReason.MISSING_IDENTIFIER
        assert match_audit_identifier(db, "A-B").reason is AuditMatchReason.NO_NORMALIZED_CANDIDATE


def test_editor_audit_import_creates_one_case_per_row_without_source_mutation(database: sessionmaker[Session], tmp_path: Path) -> None:
    add_user_and_center(database)
    with database() as db:
        db.add_all(
            [
                Asset(original_code="EXACT", normalized_code="EXACT"),
                Asset(original_code="Norm Code", normalized_code="NORMCODE"),
                Asset(original_code="dup code", normalized_code="DUPCODE"),
                Asset(original_code="DUPCODE", normalized_code="DUPCODE"),
            ]
        )
        db.commit()
    content = workbook_bytes(
        {"Physical": [
            [2, "EXACT", "machine", "brand", "model", "works", "good", "secret note", "yes", "190", "190", "DEVOLVIO"],
            [1, " Norm Code ", "machine", "brand", "model", "fails", "bad", "secret note", "yes", "190", "190", "VOLVI� 4"],
            [1, "DUP CODE", "machine", "brand", "model", "works", "good", "secret note", "yes", "190", "190", None],
            [1, None, "machine", "brand", "model", "works", "good", "secret note", "yes", "190", "190", None],
        ]},
        reference_sheet=True,
    )

    async def submit() -> httpx.Response:
        active = await client()
        try:
            return await active.post(
                "/api/imports/audit", files={"file": ("physical.xlsx", content)},
                data={"report_date": "2026-06-27", "cost_center_code": "190"},
            )
        finally:
            await active.aclose()

    response = asyncio.run(submit())
    assert response.status_code == 201
    assert response.json()["match_counts"] == {
        "exact_original_code": 1, "normalized_code": 1, "missing_identifier": 1,
        "no_normalized_candidate": 0, "ambiguous_normalized_candidate": 1,
    }
    assert response.json()["current_state_counts"] == {"found": 0, "returned": 1, "review_required": 1}
    assert "secret note" not in response.text
    with database() as db:
        observations = list(db.scalars(select(AssetObservation).where(AssetObservation.source == "audit")))
        cases = list(db.scalars(select(ReconciliationCase)))
        batch = db.scalar(select(ImportBatch).where(ImportBatch.source == "audit"))
        assert len(observations) == len(cases) == 4
        assert len(list(db.scalars(select(Asset)))) == 4
        assert list(db.scalars(select(AssetAlias))) == []
        assert all(observation.asset_id is None for observation in observations)
        assert sorted(str(case.match_strategy.value if case.match_strategy else case.match_reason.value) for case in cases) == [
            "ambiguous_normalized_candidate", "exact_original_code", "missing_identifier", "normalized_code"
        ]
        assert observations[0].quantity == 2
        assert observations[0].original_data["cells"][7]["value"] == "secret note"
        assert observations[0].original_data["cells"][11]["value"] == "DEVOLVIO"
        projections = list(db.scalars(select(AuditCurrentStateProjection)))
        assert len(projections) == 2
        exact_projection = next(projection for projection in projections if projection.marker_raw_value == "DEVOLVIO")
        assert exact_projection.state is AuditCurrentState.RETURNED
        assert exact_projection.audit_observation_id == observations[0].id
        assert (tmp_path / "imports" / batch.storage_path).read_bytes() == content
        viewer = db.scalar(select(User).where(User.username == "audit-user"))
        assert viewer is not None
        viewer.role = UserRole.VIEWER
        db.commit()

    async def read_current_states() -> httpx.Response:
        active = await client()
        try:
            return await active.get("/api/current-states", params={"cost_center_code": "190"})
        finally:
            await active.aclose()

    current_states = asyncio.run(read_current_states())
    assert current_states.status_code == 200
    returned = next(state for state in current_states.json()["states"] if state["state"] == "returned")
    assert returned["marker"] == {"column": 12, "category": "recognized_return", "raw_value": "DEVOLVIO"}
    review_required = next(state for state in current_states.json()["states"] if state["state"] == "review_required")
    assert review_required["reason"] == "ambiguous_column_l_return_marker_requires_review"
    assert review_required["marker"] == {"column": 12, "category": "ambiguous_return", "raw_value": "VOLVI� 4"}
    assert returned["projection_version"] == "audit_l_return_v1"
    assert "secret note" not in current_states.text


def test_audit_import_resolves_selected_center_before_reading_malformed_or_duplicate_content(
    database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch
) -> None:
    add_user_and_center(database)
    with database() as db:
        db.add(CostCenter(code="300", name="Legacy", status=CostCenterStatus.INACTIVE))
        db.commit()
    duplicate_content = workbook_bytes(
        {"Physical": [[1, "ID", "machine", "brand", "model", "works", "good", None, "yes", "190", "190"]]}
    )

    async def exercise() -> list[httpx.Response]:
        active = await client()
        try:
            accepted = await active.post(
                "/api/imports/audit",
                files={"file": ("accepted.xlsx", duplicate_content)},
                data={"report_date": "2026-06-27", "cost_center_code": "190"},
            )
            assert accepted.status_code == 201

            def fail_if_read(_: object) -> bytes:
                raise AssertionError("invalid selected centers must be rejected before reading workbook bytes")

            monkeypatch.setattr(imports_api, "read_upload_content", fail_if_read)
            responses = []
            for center_code in ("missing", "300"):
                for filename, content in (
                    ("malformed.xlsx", b"not an xlsx"),
                    ("duplicate.xlsx", duplicate_content),
                ):
                    responses.append(
                        await active.post(
                            "/api/imports/audit",
                            files={"file": (filename, content)},
                            data={"report_date": "2026-06-28", "cost_center_code": center_code},
                        )
                    )
            return responses
        finally:
            await active.aclose()

    responses = asyncio.run(exercise())
    assert [(response.status_code, response.json()["detail"]) for response in responses] == [
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code is inactive"),
        (422, "cost_center_code is inactive"),
    ]
    with database() as db:
        reasons = list(
            db.scalars(
                select(AuditLog.details)
                .where(AuditLog.action == "imports.audit_rejected")
                .order_by(AuditLog.created_at, AuditLog.id)
            )
        )
        assert sorted(reason["reason"] for reason in reasons) == [
            "inactive_cost_center",
            "inactive_cost_center",
            "unknown_cost_center",
            "unknown_cost_center",
        ]


def test_audit_import_requires_explicit_active_center_and_preserves_duplicate_behavior(
    database: sessionmaker[Session], tmp_path: Path
) -> None:
    add_user_and_center(database)
    with database() as db:
        db.add_all(
            [
                CostCenter(code="200", name="Finance", status=CostCenterStatus.ACTIVE),
                CostCenter(code="300", name="Legacy", status=CostCenterStatus.INACTIVE),
            ]
        )
        db.commit()
    content = workbook_bytes(
        {"Physical": [[1, "ID", "machine", "brand", "model", "works", "good", "private note", "yes", "190", "190"]]}
    )
    second_content = workbook_bytes(
        {"Physical": [[1, "SECOND", "machine", "brand", "model", "works", "good", None, "yes", "190", "190"]]}
    )

    async def exercise() -> tuple[list[httpx.Response], httpx.Response, httpx.Response, httpx.Response]:
        active = await client()
        try:
            rejected = []
            for data in (
                {"report_date": "2026-06-27"},
                {"report_date": "2026-06-27", "cost_center_code": "   "},
                {"report_date": "2026-06-27", "cost_center_code": "missing"},
                {"report_date": "2026-06-27", "cost_center_code": "300"},
            ):
                rejected.append(
                    await active.post(
                        "/api/imports/audit",
                        files={"file": ("physical.xlsx", content)},
                        data=data,
                    )
                )
            accepted = await active.post(
                "/api/imports/audit",
                files={"file": ("physical-200.xlsx", content)},
                data={"report_date": "2026-06-27", "cost_center_code": " 200 "},
            )
            accepted_second_center = await active.post(
                "/api/imports/audit",
                files={"file": ("physical-190.xlsx", second_content)},
                data={"report_date": "2026-06-28", "cost_center_code": "190"},
            )
            duplicate = await active.post(
                "/api/imports/audit",
                files={"file": ("again.xlsx", content)},
                data={"report_date": "2026-06-29", "cost_center_code": "190"},
            )
            return rejected, accepted, accepted_second_center, duplicate
        finally:
            await active.aclose()

    rejected, accepted, accepted_second_center, duplicate = asyncio.run(exercise())
    assert [(response.status_code, response.json()["detail"]) for response in rejected] == [
        (422, "cost_center_code is required"),
        (422, "cost_center_code is required"),
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code is inactive"),
    ]
    assert all("private note" not in response.text for response in rejected)
    assert accepted.status_code == 201
    assert accepted.json()["cost_center"]["code"] == "200"
    assert accepted_second_center.status_code == 201
    assert accepted_second_center.json()["cost_center"]["code"] == "190"
    assert duplicate.status_code == 409
    assert duplicate.json()["detail"] == "this file content was already imported"
    with database() as db:
        batches = list(db.scalars(select(ImportBatch).where(ImportBatch.source == "audit")))
        observations = list(db.scalars(select(AssetObservation).where(AssetObservation.source == "audit")))
        cases = list(db.scalars(select(ReconciliationCase)))
        centers = {
            center.code: center.id
            for center in db.scalars(select(CostCenter).where(CostCenter.code.in_(["190", "200"])))
        }
        expected_batch_centers = {
            "physical-190.xlsx": centers["190"],
            "physical-200.xlsx": centers["200"],
        }
        assert len(batches) == 2
        assert all(batch.cost_center_id == expected_batch_centers[batch.original_filename] for batch in batches)
        batch_center_ids = {batch.id: batch.cost_center_id for batch in batches}
        assert len(observations) == 2
        assert all(
            observation.cost_center_id == batch_center_ids[observation.import_batch_id]
            for observation in observations
        )
        observation_center_ids = {observation.id: observation.cost_center_id for observation in observations}
        assert len(cases) == 2
        assert all(case.cost_center_id == observation_center_ids[case.audit_observation_id] for case in cases)
        rejections = list(
            db.scalars(
                select(AuditLog)
                .where(AuditLog.action == "imports.audit_rejected")
                .order_by(AuditLog.created_at, AuditLog.id)
            )
        )
        assert sorted(entry.details["reason"] for entry in rejections) == sorted(
            [
                "missing_cost_center_code",
                "missing_cost_center_code",
                "unknown_cost_center",
                "inactive_cost_center",
                "duplicate_sha256",
            ]
        )
        assert len(list((tmp_path / "imports").rglob("*.xlsx"))) == 2


def test_audit_body_limit_rejects_before_multipart_parser(monkeypatch: pytest.MonkeyPatch) -> None:
    parser_calls: list[bool] = []
    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", "1")
    monkeypatch.setattr(imports_api, "parse_audit_report", lambda _: parser_calls.append(True))

    async def submit() -> httpx.Response:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as active:
            return await active.post("/api/imports/audit", content=b"xx")

    response = asyncio.run(submit())
    assert response.status_code == 413
    assert parser_calls == []


def test_audit_import_requires_editor(database: sessionmaker[Session]) -> None:
    add_user_and_center(database, UserRole.VIEWER)
    content = workbook_bytes({"Physical": [[1] + [None] * 10]})

    async def submit() -> httpx.Response:
        active = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver")
        try:
            await active.post("/api/auth/login", json={"username": "audit-user", "password": "correct horse battery staple"})
            return await active.post("/api/imports/audit", files={"file": ("physical.xlsx", content)}, data={"report_date": "2026-06-27", "cost_center_code": "190"})
        finally:
            await active.aclose()

    assert asyncio.run(submit()).status_code == 403


def test_audit_case_constraints_and_evidence_immutability_after_migration(tmp_path: Path) -> None:
    from alembic import command
    from alembic.config import Config

    config = Config(str(Path(__file__).resolve().parents[2] / "backend" / "alembic.ini"))
    database_path = tmp_path / "audit.sqlite"
    config.set_main_option("sqlalchemy.url", f"sqlite+pysqlite:///{database_path.as_posix()}")
    command.upgrade(config, "head")
    engine = create_engine(f"sqlite+pysqlite:///{database_path.as_posix()}")
    with engine.begin() as db:
        db.execute(text("INSERT INTO cost_centers (id, code, name, status) VALUES ('0000000000000000000000000000000000', '190', 'Operations', 'active')"))
        db.execute(text("INSERT INTO import_batches (id, source, cost_center_id, report_date, status, original_filename, sha256, row_count) VALUES ('0000000000000000000000000000000001', 'audit', '0000000000000000000000000000000000', '2026-06-27', 'completed', 'audit.xlsx', :sha, 1)"), {"sha": "a" * 64})
        db.execute(text("INSERT INTO asset_observations (id, import_batch_id, source, event, observed_on, original_data, quantity) VALUES ('0000000000000000000000000000000002', '0000000000000000000000000000000001', 'audit', 'snapshot', '2026-06-27', '{}', 1)"))
        db.execute(text("INSERT INTO reconciliation_cases (id, audit_observation_id, cost_center_id, match_reason) VALUES ('0000000000000000000000000000000003', '0000000000000000000000000000000002', '0000000000000000000000000000000000', 'missing_identifier')"))
        with pytest.raises(IntegrityError):
            db.execute(text("INSERT INTO reconciliation_cases (id, audit_observation_id, cost_center_id, match_strategy, match_reason) VALUES ('0000000000000000000000000000000004', '0000000000000000000000000000000002', '0000000000000000000000000000000000', 'normalized_code', 'missing_identifier')"))
        with pytest.raises(IntegrityError, match="immutable"):
            db.execute(text("UPDATE asset_observations SET quantity = 2 WHERE id = '0000000000000000000000000000000002'"))
    engine.dispose()
