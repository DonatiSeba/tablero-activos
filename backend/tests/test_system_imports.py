import asyncio
from collections.abc import Generator
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app import db as db_api
from backend.app import imports as imports_api
from backend.app import request_limits as request_limits_api
from backend.app import system_imports
from backend.app.auth import password_hasher
from backend.app.db import Base, get_db
from backend.app.main import app
from backend.app.models import (
    Asset,
    AssetObservation,
    AuditLog,
    CostCenter,
    CostCenterStatus,
    ImportBatch,
    User,
    UserRole,
)
from backend.app.system_imports import (
    ImportValidationError,
    StorageIntegrityError,
    normalize_asset_code,
    parse_system_report,
    sha256_hex,
    store_original_file,
)

HEADERS = [
    "Rubro", "Categoría", "Producto", "Cód. Ident.", "Identificación", "Nro. CC", "Centro de Costo", "Estado",
]


def workbook_bytes(rows: list[list[object]], headers: list[str] = HEADERS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def database(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Generator[sessionmaker[Session], None, None]:
    monkeypatch.setenv("SESSION_SECRET", "test-session-secret-that-is-long-enough-for-signing")
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("SESSION_TTL_SECONDS", "3600")
    monkeypatch.setenv("IMPORT_STORAGE_PATH", str(tmp_path / "imports"))
    engine = __import__("sqlalchemy").create_engine(
        "sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def override_get_db() -> Generator[Session, None, None]:
        db = factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    yield factory
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()


@pytest.fixture
def oversized_audit_database(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> Generator[sessionmaker[Session], None, None]:
    database_url = f"sqlite+pysqlite:///{tmp_path / 'oversized-audit.db'}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    db_api._session_factory.cache_clear()
    engine = __import__("sqlalchemy").create_engine(database_url)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    try:
        yield factory
    finally:
        audit_engine = db_api._session_factory().kw["bind"]
        db_api._session_factory.cache_clear()
        audit_engine.dispose()
        Base.metadata.drop_all(engine)
        engine.dispose()


def add_user(factory: sessionmaker[Session], role: UserRole) -> User:
    user = User(
        username=f"{role.value}-user", email=f"{role.value}@example.test", display_name="Import User",
        password_hash=password_hasher.hash("correct horse battery staple"), role=role, is_active=True,
    )
    with factory() as db:
        db.add_all(
            [
                user,
                CostCenter(code="190", name="Operations", status=CostCenterStatus.ACTIVE),
            ]
        )
        db.commit()
    return user


async def authenticated_client(role: UserRole) -> tuple[httpx.AsyncClient, httpx.Response]:
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver")
    login = await client.post("/api/auth/login", json={"username": f"{role.value}-user", "password": "correct horse battery staple"})
    return client, login


def report_rows() -> list[list[object]]:
    return [
        ["IT", "Laptop", "ThinkPad", "ASSET-01", "possibly-truncated", 190, "Operations", "Active"],
        ["IT", "Laptop", "ThinkPad", "ASSET-02", "possibly-truncated", 190, "OPERATIONS", "Retired"],
    ]


def test_parser_preserves_all_columns_and_validates_named_columns() -> None:
    report = parse_system_report(workbook_bytes(report_rows()))
    assert report.cost_center_code == "190"
    assert report.observed_cost_center_names == ("OPERATIONS", "Operations")
    assert report.rows[0].named_values["Cód. Ident."] == "ASSET-01"
    assert normalize_asset_code(" asset- 01 ") == "ASSET-01"
    extra = parse_system_report(workbook_bytes([report_rows()[0] + ["Supplier", 42.5]], HEADERS + ["Proveedor", "Importe"]))
    assert extra.rows[0].original_data["cells"][-2:] == [
        {"column": 9, "header": "Proveedor", "value": "Supplier"},
        {"column": 10, "header": "Importe", "value": 42.5},
    ]
    adversarial_headers = HEADERS + ["", "", "[column 11]", "Proveedor", "Proveedor"]
    adversarial = parse_system_report(
        workbook_bytes([report_rows()[0] + ["first blank", "second blank", "generated", "first supplier", "second supplier"]], adversarial_headers)
    )
    assert adversarial.rows[0].original_data["cells"][-5:] == [
        {"column": 9, "header": None, "value": "first blank"},
        {"column": 10, "header": None, "value": "second blank"},
        {"column": 11, "header": "[column 11]", "value": "generated"},
        {"column": 12, "header": "Proveedor", "value": "first supplier"},
        {"column": 13, "header": "Proveedor", "value": "second supplier"},
    ]
    with pytest.raises(ImportValidationError, match="required column") as missing:
        parse_system_report(workbook_bytes(report_rows(), HEADERS[:-1]))
    assert missing.value.code == "missing_required_columns"
    with pytest.raises(ImportValidationError, match="readable") as unreadable:
        parse_system_report(b"not an xlsx")
    assert unreadable.value.code == "unreadable_workbook"


def test_storage_publish_is_atomic_and_existing_content_is_verified(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    root = tmp_path / "imports"
    monkeypatch.setenv("IMPORT_STORAGE_PATH", str(root))
    content = b"immutable evidence"
    digest = sha256_hex(content)
    with ThreadPoolExecutor(max_workers=8) as executor:
        results = list(executor.map(lambda _: store_original_file(content, digest), range(8)))
    assert sum(created for _, created in results) == 1
    target = root / results[0][0]
    assert target.read_bytes() == content
    assert not list(target.parent.glob(".*.tmp"))
    target.write_bytes(b"corrupt")
    with pytest.raises(StorageIntegrityError):
        store_original_file(content, digest)


def test_editor_import_preserves_extra_evidence_and_raw_file(database: sessionmaker[Session], tmp_path: Path) -> None:
    add_user(database, UserRole.EDITOR)
    content = workbook_bytes([report_rows()[0] + ["Supplier A"]], HEADERS + ["Proveedor"])

    async def exercise() -> httpx.Response:
        client, login = await authenticated_client(UserRole.EDITOR)
        assert login.status_code == 200
        try:
            return await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", content)},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
        finally:
            await client.aclose()

    response = asyncio.run(exercise())
    assert response.status_code == 201
    with database() as db:
        batch = db.scalar(select(ImportBatch))
        observation = db.scalar(select(AssetObservation))
        assert batch is not None and observation is not None
        assert (tmp_path / "imports" / batch.storage_path).read_bytes() == content
        assert observation.original_data["cells"][-1] == {"column": 9, "header": "Proveedor", "value": "Supplier A"}
        assert db.scalar(select(AuditLog.action).where(AuditLog.action == "imports.system_accepted")) is not None


def test_duplicate_and_invalid_imports_are_rejected_and_audited(database: sessionmaker[Session]) -> None:
    add_user(database, UserRole.ADMIN)
    content = workbook_bytes(report_rows())

    async def exercise() -> tuple[httpx.Response, httpx.Response]:
        client, _ = await authenticated_client(UserRole.ADMIN)
        try:
            data = {"report_date": "2026-09-17", "cost_center_code": "190"}
            first = await client.post("/api/imports/system", files={"file": ("cc190.xlsx", content)}, data=data)
            duplicate = await client.post("/api/imports/system", files={"file": ("again.xlsx", content)}, data=data)
            return first, duplicate
        finally:
            await client.aclose()

    first, duplicate = asyncio.run(exercise())
    assert first.status_code == 201
    assert duplicate.status_code == 409 and duplicate.json()["detail"] == "this file content was already imported"
    with database() as db:
        assert len(list(db.scalars(select(ImportBatch)))) == 1
        assert db.scalar(select(AuditLog).where(AuditLog.action == "imports.system_rejected")).details == {"reason": "duplicate_sha256"}


def test_system_import_resolves_selected_center_before_reading_malformed_or_duplicate_content(
    database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch
) -> None:
    add_user(database, UserRole.EDITOR)
    with database() as db:
        db.add(CostCenter(code="300", name="Legacy", status=CostCenterStatus.INACTIVE))
        db.commit()
    duplicate_content = workbook_bytes(report_rows())

    async def exercise() -> list[httpx.Response]:
        client, _ = await authenticated_client(UserRole.EDITOR)
        try:
            accepted = await client.post(
                "/api/imports/system",
                files={"file": ("accepted.xlsx", duplicate_content)},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
            assert accepted.status_code == 201

            def fail_if_read(_: object) -> bytes:
                raise AssertionError("invalid selected centers must be rejected before reading workbook bytes")

            monkeypatch.setattr(imports_api, "read_upload_content", fail_if_read)
            responses = []
            for center_code in ("missing", "300"):
                for filename, content in (
                    ("malformed.xlsx", b"not an xlsx"),
                    ("duplicate.xlsx", duplicate_content),
                ):
                    responses.append(
                        await client.post(
                            "/api/imports/system",
                            files={"file": (filename, content)},
                            data={"report_date": "2026-09-18", "cost_center_code": center_code},
                        )
                    )
            return responses
        finally:
            await client.aclose()

    responses = asyncio.run(exercise())
    assert [(response.status_code, response.json()["detail"]) for response in responses] == [
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code is inactive"),
        (422, "cost_center_code is inactive"),
    ]
    with database() as db:
        reasons = list(
            db.scalars(
                select(AuditLog.details)
                .where(AuditLog.action == "imports.system_rejected")
                .order_by(AuditLog.created_at, AuditLog.id)
            )
        )
        assert sorted(reason["reason"] for reason in reasons) == [
            "inactive_cost_center",
            "inactive_cost_center",
            "unknown_cost_center",
            "unknown_cost_center",
        ]


def test_system_import_requires_explicit_active_matching_center_without_persisting_rejected_evidence(
    database: sessionmaker[Session], tmp_path: Path
) -> None:
    add_user(database, UserRole.EDITOR)
    with database() as db:
        db.add_all(
            [
                CostCenter(code="200", name="Finance", status=CostCenterStatus.ACTIVE),
                CostCenter(code="300", name="Legacy", status=CostCenterStatus.INACTIVE),
            ]
        )
        db.commit()
    content = workbook_bytes(report_rows())

    async def exercise() -> tuple[list[httpx.Response], httpx.Response, httpx.Response]:
        client, _ = await authenticated_client(UserRole.EDITOR)
        try:
            rejected = []
            for data in (
                {"report_date": "2026-09-17"},
                {"report_date": "2026-09-17", "cost_center_code": "   "},
                {"report_date": "2026-09-17", "cost_center_code": "missing"},
                {"report_date": "2026-09-17", "cost_center_code": "300"},
                {"report_date": "2026-09-17", "cost_center_code": "200"},
            ):
                rejected.append(
                    await client.post(
                        "/api/imports/system",
                        files={"file": ("cc190.xlsx", content)},
                        data=data,
                    )
                )
            accepted = await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", content)},
                data={"report_date": "2026-09-17", "cost_center_code": " 190 "},
            )
            rows_200 = [row[:5] + [200, "Finance"] + row[7:] for row in report_rows()]
            accepted_second_center = await client.post(
                "/api/imports/system",
                files={"file": ("cc200.xlsx", workbook_bytes(rows_200))},
                data={"report_date": "2026-09-18", "cost_center_code": "200"},
            )
            return rejected, accepted, accepted_second_center
        finally:
            await client.aclose()

    rejected, accepted, accepted_second_center = asyncio.run(exercise())
    assert [(response.status_code, response.json()["detail"]) for response in rejected] == [
        (422, "cost_center_code is required"),
        (422, "cost_center_code is required"),
        (422, "cost_center_code does not exist"),
        (422, "cost_center_code is inactive"),
        (422, "workbook Nro. CC does not match cost_center_code"),
    ]
    assert accepted.status_code == 201
    assert accepted.json()["cost_center"]["code"] == "190"
    assert accepted_second_center.status_code == 201
    assert accepted_second_center.json()["cost_center"]["code"] == "200"
    with database() as db:
        batches = list(db.scalars(select(ImportBatch)))
        observations = list(db.scalars(select(AssetObservation)))
        assets = list(db.scalars(select(Asset)))
        centers = list(db.scalars(select(CostCenter).order_by(CostCenter.code)))
        rejections = list(
            db.scalars(
                select(AuditLog)
                .where(AuditLog.action == "imports.system_rejected")
                .order_by(AuditLog.created_at, AuditLog.id)
            )
        )
        center_ids = {center.code: center.id for center in centers}
        expected_batch_centers = {
            "cc190.xlsx": center_ids["190"],
            "cc200.xlsx": center_ids["200"],
        }
        assert len(batches) == 2
        assert all(batch.cost_center_id == expected_batch_centers[batch.original_filename] for batch in batches)
        batch_center_ids = {batch.id: batch.cost_center_id for batch in batches}
        assert len(observations) == 4 and len(assets) == 2
        assert all(
            observation.cost_center_id == batch_center_ids[observation.import_batch_id]
            for observation in observations
        )
        assert [center.code for center in centers] == ["190", "200", "300"]
        assert sorted(entry.details["reason"] for entry in rejections) == sorted(
            [
                "missing_cost_center_code",
                "missing_cost_center_code",
                "unknown_cost_center",
                "inactive_cost_center",
                "cost_center_mismatch",
            ]
        )
        assert len(list((tmp_path / "imports").rglob("*.xlsx"))) == 2


def test_lazy_workbook_failure_is_audited_validation_rejection(database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch) -> None:
    add_user(database, UserRole.EDITOR)

    class LazyFailureWorkbook:
        @property
        def active(self):
            raise RuntimeError("lazy reader failure")

        def close(self) -> None:
            pass

    monkeypatch.setattr(system_imports, "load_workbook", lambda **_: LazyFailureWorkbook())

    async def exercise() -> httpx.Response:
        client, _ = await authenticated_client(UserRole.EDITOR)
        try:
            return await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", workbook_bytes(report_rows()))},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
        finally:
            await client.aclose()

    response = asyncio.run(exercise())
    assert response.status_code == 422
    assert response.json()["detail"] == "file is not a readable .xlsx workbook"
    with database() as db:
        assert db.scalar(select(AuditLog).where(AuditLog.action == "imports.system_rejected")).details == {"reason": "unreadable_workbook"}


def test_nonduplicate_integrity_error_is_audited_and_keeps_digest_file(
    database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    add_user(database, UserRole.EDITOR)
    content = workbook_bytes(report_rows())

    def fail_asset(*_: object, **__: object) -> Asset:
        raise IntegrityError("insert", {}, RuntimeError("unrelated constraint"))

    monkeypatch.setattr(imports_api, "_asset_for_exact_code", fail_asset)

    async def exercise() -> httpx.Response:
        client, _ = await authenticated_client(UserRole.EDITOR)
        try:
            return await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", content)},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
        finally:
            await client.aclose()

    response = asyncio.run(exercise())
    assert response.status_code == 409
    assert response.json()["detail"] == "could not persist import evidence because of a database constraint"
    digest = sha256_hex(content)
    assert (tmp_path / "imports" / digest[:2] / f"{digest}.xlsx").read_bytes() == content
    with database() as db:
        rejected = db.scalar(select(AuditLog).where(AuditLog.action == "imports.system_rejected"))
        assert rejected is not None and rejected.details == {"reason": "database_integrity_error"}


def test_in_handler_upload_bound_remains_a_defense_in_depth(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", "3")
    with pytest.raises(ImportValidationError, match="uploaded file exceeds") as oversized:
        system_imports.read_upload_content(BytesIO(b"1234"))
    assert oversized.value.code == "upload_too_large"


def test_upload_and_archive_limits_reject_without_evidence_storage(database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    add_user(database, UserRole.EDITOR)
    content = workbook_bytes(report_rows())

    async def submit() -> httpx.Response:
        client, _ = await authenticated_client(UserRole.EDITOR)
        try:
            return await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", content)},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
        finally:
            await client.aclose()

    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", "1")
    too_large = asyncio.run(submit())
    assert too_large.status_code == 413 and too_large.json()["detail"] == "request body exceeds the configured upload size limit"
    assert not (tmp_path / "imports").exists()
    # Leave enough room for multipart framing so the in-handler XLSX archive bound is exercised.
    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", str(len(content) + 1024))
    monkeypatch.setenv("IMPORT_MAX_XLSX_UNCOMPRESSED_BYTES", "1")
    expanded = asyncio.run(submit())
    assert expanded.status_code == 422 and expanded.json()["detail"] == "workbook exceeds the configured uncompressed size limit"
    assert not (tmp_path / "imports").exists()


def test_oversized_declared_and_streamed_bodies_bypass_import_handler_and_write_sanitized_audits(
    oversized_audit_database: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch
) -> None:
    handler_calls: list[object] = []

    def handler_should_not_run(*_: object) -> bytes:
        handler_calls.append(True)
        raise AssertionError("the import handler must not read an oversized body")

    monkeypatch.setattr(imports_api, "read_upload_content", handler_should_not_run)
    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", "256")
    declared_content = workbook_bytes(report_rows())
    boundary = "test-boundary"
    streamed_body = (
        f"--{boundary}\r\n"
        "Content-Disposition: form-data; name=\"file\"; filename=\"cc190.xlsx\"\r\n"
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + (b"x" * 512) + f"\r\n--{boundary}--\r\n".encode()

    async def stream_body():
        yield streamed_body[:128]
        yield streamed_body[128:]

    async def exercise() -> tuple[httpx.Response, httpx.Response]:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
            declared = await client.post(
                "/api/imports/system",
                files={"file": ("cc190.xlsx", declared_content)},
                data={"report_date": "2026-09-17", "cost_center_code": "190"},
            )
            streamed = await client.post(
                "/api/imports/system",
                content=stream_body(),
                headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
            )
            return declared, streamed

    declared, streamed = asyncio.run(exercise())
    assert declared.status_code == 413
    assert streamed.status_code == 413
    assert handler_calls == []
    with oversized_audit_database() as db:
        rejections = list(
            db.scalars(select(AuditLog).where(AuditLog.action == "imports.system_rejected").order_by(AuditLog.created_at))
        )
    assert len(rejections) == 2
    assert all(rejection.target_entity == "import" for rejection in rejections)
    assert all(rejection.user_id is None for rejection in rejections)
    assert all(rejection.details == {"reason": "request_too_large"} for rejection in rejections)
    assert all(rejection.ip_address is not None for rejection in rejections)


def test_oversized_body_stays_413_when_audit_persistence_is_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("IMPORT_MAX_UPLOAD_BYTES", "1")

    def unavailable_session_factory() -> object:
        raise RuntimeError("audit database unavailable")

    monkeypatch.setattr(request_limits_api, "_session_factory", unavailable_session_factory)

    async def exercise() -> httpx.Response:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
            return await client.post("/api/imports/system", content=b"xx")

    response = asyncio.run(exercise())
    assert response.status_code == 413


def test_system_import_requires_editor_role_and_audits_denial(database: sessionmaker[Session]) -> None:
    add_user(database, UserRole.VIEWER)

    async def exercise() -> httpx.Response:
        client, login = await authenticated_client(UserRole.VIEWER)
        assert login.status_code == 200
        try:
            return await client.post("/api/imports/system", files={"file": ("cc190.xlsx", workbook_bytes(report_rows()))}, data={"report_date": "2026-09-17"})
        finally:
            await client.aclose()

    response = asyncio.run(exercise())
    assert response.status_code == 403
    with database() as db:
        denial = db.scalar(select(AuditLog).where(AuditLog.action == "auth.authorization_denied"))
        assert denial is not None and denial.details == {"required_role": "editor", "actual_role": "viewer"}
